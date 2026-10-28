"""
Bulk import of Sections and Technical Data Sheet / Pricing Tool fields from one Excel
workbook — the Formulas page's "Add Section" and "Add Field" flows, many rows at once.

GET  /api/formulas/bulk-import/template/ — the blank workbook to fill in
POST /api/formulas/bulk-import/          — multipart `file` (.xlsx), plus `dry_run`
                                           ("true" by default: validate and report only)

Workbook layout (sheet and column names are case-insensitive):
  "Sections": label | sheet | order
  "Fields":   label | key | section | type | unit | fixed_value | auto | expression

Every row goes through the same checks as the single-row endpoints (section must be
registered, unique keys, unknown variables, expression dry-run, Text/Select Auto fields
copy-only, fixed value one of the options / numeric). A section or field that already
exists is skipped, not changed. Rows are applied inside one transaction: a dry run, or a
real import with any error, rolls everything back — nothing is ever half-imported.
"""

import io
import re

from django.db import transaction
from django.db.models import Max
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import status
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from django.http import HttpResponse

from accounts.permissions import CanEditReferenceData

from .evaluator import extract_variables
from .models import TECH_DATA_AUTO_SECTION_LABEL, FormulaDefinition, Section, SectionTab, TechDataFieldDefinition
from .serializers import TechDataFieldCreateSerializer
from .views import resolve_expression_variables, validate_auto_field_type

MAX_FILE_SIZE_BYTES = 2 * 1024 * 1024
MAX_ROWS_PER_SHEET = 1000

SECTION_COLUMNS = ["label", "sheet", "order"]
# No `options` column: a Manual select field is imported without choices, which are
# then added (or a catalog linked) with Edit on the Formulas page.
FIELD_COLUMNS = ["label", "key", "section", "type", "unit", "fixed_value", "auto", "expression"]
SELECT_WITHOUT_OPTIONS_NOTE = "Dropdown imported without choices — add them with Edit on the Formulas page."
SELECT_WITH_FIXED_VALUE_NOTE = "Dropdown imported with one choice, '{}' (pre-filled) — add more with Edit on the Formulas page."

EXAMPLE_SECTIONS = [
    ["9. Coupling Parts", "pricing_tool", 24],
    ["Extra Inputs", "tech_sheet", None],
]
EXAMPLE_FIELDS = [
    ["Coupling Weight", "couplingWeight", "9. Coupling Parts", "number", "kg", None, "No", None],
    ["Coupling Type", None, "9. Coupling Parts", "select", None, None, "No", None],
    ["Coupling Cost", None, "9. Coupling Parts", "number", "INR", None, "Yes", "couplingWeight * 120"],
]


class BulkFieldCreateSerializer(TechDataFieldCreateSerializer):
    """TechDataFieldCreateSerializer without its "a Manual select needs at least one
    option" rule — the workbook has no options column (see FIELD_COLUMNS)."""

    def validate(self, data):
        if data["is_auto"] and not data.get("expression", "").strip():
            raise ValidationError({"expression": "Required for an Auto field."})
        return data

SHEET_ALIASES = {
    "pricing tool": SectionTab.PRICING_TOOL,
    "pricing tool formulas": SectionTab.PRICING_TOOL,
    "pricing": SectionTab.PRICING_TOOL,
    "tech sheet": SectionTab.TECH_SHEET,
    "technical sheet": SectionTab.TECH_SHEET,
    "technical sheet formulas": SectionTab.TECH_SHEET,
    "technical data sheet": SectionTab.TECH_SHEET,
    "tech": SectionTab.TECH_SHEET,
}
YES = {"yes", "y", "true", "1", "auto"}
NO = {"", "no", "n", "false", "0", "manual"}


def slugify_to_camel_case(label: str) -> str:
    """Same key the Formulas page's forms derive from a label (FormulasPage.tsx's
    slugifyToCamelCase), so a key left blank in the workbook matches what Add Field
    would have suggested."""
    words = [w for w in re.split(r"[^a-zA-Z0-9]+", label.strip()) if w]
    return "".join(w[0].lower() + w[1:] if i == 0 else w[0].upper() + w[1:] for i, w in enumerate(words))


def cell_text(value) -> str:
    """A cell as the text the user typed — Excel hands back 24 as 24.0."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def error_text(exc: ValidationError) -> str:
    detail = exc.detail
    if isinstance(detail, dict):
        parts = []
        for field, messages in detail.items():
            messages = messages if isinstance(messages, list) else [messages]
            text = " ".join(str(m) for m in messages)
            parts.append(text if field in ("detail", "non_field_errors") else f"{field}: {text}")
        return "; ".join(parts)
    if isinstance(detail, list):
        return " ".join(str(m) for m in detail)
    return str(detail)


def read_sheet(workbook, name: str, columns: list[str]) -> list[dict] | None:
    """The rows of sheet `name` as dicts keyed by `columns` (plus `row`, the Excel row
    number), skipping blank rows. None when the workbook has no such sheet."""
    sheet = next((ws for ws in workbook.worksheets if ws.title.strip().lower() == name.lower()), None)
    if sheet is None:
        return None
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    positions = {}
    for i, title in enumerate(header):
        normalized = re.sub(r"[\s-]+", "_", cell_text(title).lower())
        if normalized in columns:
            positions[normalized] = i
    if "label" not in positions:
        raise ValidationError({"detail": f"Sheet '{sheet.title}' needs a header row with at least a 'label' column."})

    result = []
    for row_number, values in enumerate(rows, start=2):
        record = {col: cell_text(values[i]) if i < len(values) else "" for col, i in positions.items()}
        if not any(record.values()):
            continue
        record.update({col: "" for col in columns if col not in record})
        record["row"] = row_number
        result.append(record)
        if len(result) > MAX_ROWS_PER_SHEET:
            raise ValidationError({"detail": f"Sheet '{sheet.title}' has more than {MAX_ROWS_PER_SHEET} rows — split it into smaller files."})
    return result


class BulkImporter:
    """Applies one workbook's rows; the caller owns the outer transaction."""

    def __init__(self, user):
        self.user = user
        self.errors: list[dict] = []
        self.sections: list[dict] = []
        self.fields: list[dict] = []

    def fail(self, sheet: str, row: int, message: str, report: dict | None = None) -> None:
        self.errors.append({"sheet": sheet, "row": row, "message": message})
        if report is not None:
            report.update(status="error", message=message)

    # --- Sections ---------------------------------------------------------------

    def import_sections(self, rows: list[dict]) -> None:
        next_order = (Section.objects.aggregate(m=Max("order"))["m"] or 0) + 1
        used_keys = set(Section.objects.values_list("key", flat=True))
        seen_labels = set()

        for row in rows:
            label = row["label"]
            report = {"row": row["row"], "label": label, "sheet": None, "order": None, "key": None, "status": "", "message": ""}
            self.sections.append(report)
            if not label:
                self.fail("Sections", row["row"], "label is required.", report)
                continue
            if len(label) > 200:
                self.fail("Sections", row["row"], "label is longer than 200 characters.", report)
                continue
            if label in seen_labels:
                self.fail("Sections", row["row"], f"'{label}' appears more than once in this file.", report)
                continue
            seen_labels.add(label)

            sheet_raw = row["sheet"].strip().lower()
            sheet = sheet_raw if sheet_raw in SectionTab.values else SHEET_ALIASES.get(re.sub(r"[^a-z]+", " ", sheet_raw).strip())
            if sheet is None:
                self.fail("Sections", row["row"], "sheet must be 'pricing_tool' or 'tech_sheet'.", report)
                continue
            report["sheet"] = sheet

            if Section.objects.filter(label=label).exists():
                report.update(status="skipped", message="Already exists — left unchanged.")
                continue

            if row["order"]:
                try:
                    order = int(float(row["order"]))
                    if order < 0 or order != float(row["order"]):
                        raise ValueError
                except ValueError:
                    self.fail("Sections", row["row"], "order must be a whole number of 0 or more.", report)
                    continue
            else:
                order = next_order
                next_order += 1

            base_key = (slugify_to_camel_case(label) or "section")[:44]
            key, n = base_key, 2
            while key in used_keys:
                key, n = f"{base_key}-{n}", n + 1
            used_keys.add(key)

            Section.objects.create(key=key, label=label, order=order, tab=sheet)
            report.update(key=key, order=order, status="created")

    # --- Fields -----------------------------------------------------------------

    def import_fields(self, rows: list[dict]) -> None:
        prepared = []
        seen_keys = {}
        for row in rows:
            key = row["key"] or slugify_to_camel_case(row["label"])
            report = {
                "row": row["row"], "key": key, "label": row["label"], "section": row["section"],
                "type": (row["type"] or "text").lower(), "auto": False, "status": "", "message": "",
            }
            self.fields.append(report)
            auto_raw = row["auto"].strip().lower()
            if auto_raw not in YES | NO:
                self.fail("Fields", row["row"], "auto must be Yes or No.", report)
                continue
            report["auto"] = auto_raw in YES
            if not key:
                self.fail("Fields", row["row"], "label (or key) is required.", report)
                continue
            if key in seen_keys:
                self.fail("Fields", row["row"], f"key '{key}' is also used on row {seen_keys[key]} of this file.", report)
                continue
            seen_keys[key] = row["row"]
            if TechDataFieldDefinition.objects.filter(key=key).exists():
                report.update(status="skipped", message="A field with this key already exists — left unchanged.")
                continue
            prepared.append((row, report))

        # Manual fields first, then Auto fields so each one's formula only runs once
        # every field it reads from this same file already exists.
        new_keys = {report["key"]: (row, report) for row, report in prepared}
        deps = {
            report["key"]: {v for v in extract_variables(row["expression"]) if v in new_keys and v != report["key"]}
            if report["auto"] else set()
            for row, report in prepared
        }
        done: set[str] = set()
        failed: set[str] = set()
        remaining = [report["key"] for _, report in prepared]
        while remaining:
            ready = [k for k in remaining if deps[k] <= done | failed]
            if not ready:
                for k in remaining:
                    row, report = new_keys[k]
                    self.fail("Fields", row["row"], "its expression is part of a circular reference with other fields in this file.", report)
                break
            for k in ready:
                remaining.remove(k)
                row, report = new_keys[k]
                broken = sorted(deps[k] & failed)
                if broken:
                    rows_ = ", ".join(str(new_keys[b][0]["row"]) for b in broken)
                    self.fail("Fields", row["row"], f"uses {', '.join(broken)}, which has an error (row {rows_}).", report)
                    failed.add(k)
                    continue
                try:
                    with transaction.atomic():
                        self.create_field(row, report)
                    done.add(k)
                except ValidationError as exc:
                    self.fail("Fields", row["row"], error_text(exc), report)
                    failed.add(k)

    def create_field(self, row: dict, report: dict) -> None:
        """Same rules as TechDataFieldDefinitionViewSet.create(), plus the fixed-value
        checks its perform_update() applies — the single-row form saves a fixed value
        through a follow-up PATCH, here it's part of the one create."""
        is_auto = report["auto"]
        fixed_value = row["fixed_value"]
        # A Manual select's fixed value becomes its one choice, so it's pre-filled and
        # valid; more choices (or a catalog link) are added later with Edit.
        options = [fixed_value] if report["type"] == "select" and not is_auto and fixed_value else []
        req = BulkFieldCreateSerializer(data={
            "key": report["key"], "label": row["label"], "section": row["section"], "unit": row["unit"],
            "field_type": report["type"], "options": options, "is_auto": is_auto,
            "expression": row["expression"], "output_unit": row["unit"][:20],
        })
        req.is_valid(raise_exception=True)
        data = req.validated_data

        input_variables = []
        if is_auto:
            if FormulaDefinition.objects.filter(key=data["key"]).exists():
                raise ValidationError({"key": f"A formula with the key '{data['key']}' already exists."})
            validate_auto_field_type(data["field_type"], data["expression"])
            input_variables = resolve_expression_variables(data["expression"], self_key=data["key"])

        if fixed_value:
            if is_auto:
                raise ValidationError({"fixed_value": "An Auto field can't have a fixed value."})
            if data["field_type"] == "number":
                try:
                    float(fixed_value)
                except ValueError:
                    raise ValidationError({"fixed_value": "Must be a number for a Number field."})
        if len(fixed_value) > 200:
            raise ValidationError({"fixed_value": "Longer than 200 characters."})

        order = (TechDataFieldDefinition.objects.filter(section=data["section"]).aggregate(m=Max("order"))["m"] or 0) + 1
        formula = None
        if is_auto:
            formula = FormulaDefinition.objects.create(
                key=data["key"],
                label=data["label"],
                section=TECH_DATA_AUTO_SECTION_LABEL,
                expression=data["expression"],
                input_variables=input_variables,
                output_unit=data["output_unit"],
                updated_by=self.user,
            )
        TechDataFieldDefinition.objects.create(
            key=data["key"],
            label=data["label"],
            section=data["section"],
            unit=data["unit"],
            field_type=data["field_type"],
            options=data["options"],
            fixed_value=fixed_value,
            formula=formula,
            order=order,
            is_core=False,
            created_by=self.user,
        )
        note = ""
        if data["field_type"] == "select" and not is_auto:
            note = SELECT_WITH_FIXED_VALUE_NOTE.format(fixed_value) if fixed_value else SELECT_WITHOUT_OPTIONS_NOTE
        report.update(section=data["section"], status="created", message=note)


class BulkImportView(APIView):
    permission_classes = [CanEditReferenceData]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        uploaded = request.FILES.get("file")
        if not uploaded:
            raise ValidationError({"detail": "Attach the filled-in Excel file as 'file'."})
        if not uploaded.name.lower().endswith(".xlsx"):
            raise ValidationError({"detail": "Upload an Excel .xlsx file (download the template to start)."})
        if uploaded.size > MAX_FILE_SIZE_BYTES:
            raise ValidationError({"detail": "The file is larger than 2 MB."})
        dry_run = str(request.data.get("dry_run", "true")).strip().lower() not in ("false", "0", "no")

        try:
            workbook = load_workbook(io.BytesIO(uploaded.read()), read_only=True, data_only=True)
        except (InvalidFileException, OSError, KeyError, ValueError):
            raise ValidationError({"detail": "Couldn't read this file — save it as an Excel .xlsx workbook and try again."})
        try:
            section_rows = read_sheet(workbook, "Sections", SECTION_COLUMNS)
            field_rows = read_sheet(workbook, "Fields", FIELD_COLUMNS)
        finally:
            workbook.close()
        if section_rows is None and field_rows is None:
            raise ValidationError({"detail": "The workbook needs a 'Sections' sheet, a 'Fields' sheet, or both."})

        importer = BulkImporter(request.user)
        with transaction.atomic():
            importer.import_sections(section_rows or [])
            importer.import_fields(field_rows or [])
            saved = not dry_run and not importer.errors
            if not saved:
                transaction.set_rollback(True)

        def count(rows, state):
            return sum(1 for r in rows if r["status"] == state)

        body = {
            "dry_run": dry_run,
            "saved": saved,
            "summary": {
                "sections_created": count(importer.sections, "created"),
                "sections_skipped": count(importer.sections, "skipped"),
                "fields_created": count(importer.fields, "created"),
                "fields_skipped": count(importer.fields, "skipped"),
                "errors": len(importer.errors),
            },
            "sections": importer.sections,
            "fields": importer.fields,
            "errors": importer.errors,
        }
        if not dry_run and importer.errors:
            n = len(importer.errors)
            body["detail"] = f"Nothing was imported — {n} row{'s' if n != 1 else ''} need{'' if n != 1 else 's'} fixing first."
            return Response(body, status=status.HTTP_400_BAD_REQUEST)
        return Response(body)


class BulkImportTemplateView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        workbook = Workbook()
        bold = Font(bold=True)
        sections = workbook.active
        sections.title = "Sections"
        fields = workbook.create_sheet("Fields")
        # Example rows (highlighted) show how each sheet is filled in — the preview lists
        # them as "Will be added", so they're replaced or deleted before confirming.
        example_fill = PatternFill("solid", start_color="FFF4CC")
        for sheet, columns, examples in (
            (sections, SECTION_COLUMNS, EXAMPLE_SECTIONS),
            (fields, FIELD_COLUMNS, EXAMPLE_FIELDS),
        ):
            sheet.append(columns)
            for cell in sheet[1]:
                cell.font = bold
            for example in examples:
                sheet.append(example)
                for cell in sheet[sheet.max_row]:
                    cell.fill = example_fill
            sheet.freeze_panes = "A2"
            for i in range(len(columns)):
                sheet.column_dimensions[chr(ord("A") + i)].width = 22

        guide = workbook.create_sheet("Instructions")
        guide.column_dimensions["A"].width = 16
        guide.column_dimensions["B"].width = 100
        lines = [
            ("Start here", ""),
            ("", "The yellow rows on the Sections and Fields sheets are examples — replace them with your own rows, or delete them."),
            ("", "Sections sheet = Add Section. Fields sheet = Add Field. Use either sheet or both in one file."),
            ("", ""),
            ("Sections sheet", ""),
            ("label", "Section name, e.g. 9. Coupling Parts. A section that already exists is skipped."),
            ("sheet", "pricing_tool or tech_sheet — which tab of the Formulas page it belongs to."),
            ("order", "Optional whole number. Empty = after the last section."),
            ("", ""),
            ("Fields sheet", ""),
            ("label", "Field name, e.g. Coupling Weight."),
            ("key", "Optional. Empty = made from the label (Coupling Weight -> couplingWeight). A field whose key already exists is skipped."),
            ("section", "An existing section, or one from the Sections sheet of this file."),
            ("type", "text, number or select (dropdown). Empty = text. A select field's choices are added later with Edit on the Formulas page (its fixed_value, if any, is imported as the first choice)."),
            ("unit", "Optional, e.g. kg, INR, mm."),
            ("fixed_value", "Optional, Manual fields only. Pre-filled on every new RFQ. For select it becomes the dropdown's first choice."),
            ("auto", "Yes = calculated from the expression (read-only on the sheet). No or empty = entered on each RFQ."),
            ("expression", "Auto fields only, e.g. couplingWeight * 120. Text/select Auto fields can only copy one field: just its key."),
            ("", ""),
            ("Field examples", ""),
            ("Manual number", "Coupling Weight — type number, auto No: entered on each RFQ."),
            ("Manual select", "Coupling Type — type select, auto No. After importing, click Edit on it to add its choices (e.g. Rigid, Flexible)."),
            ("Auto calculation", "Coupling Cost — type number, auto Yes, expression couplingWeight * 120."),
        ]
        for line in lines:
            guide.append(line)
            if line[0] and not line[1]:
                guide.cell(row=guide.max_row, column=1).font = bold

        buffer = io.BytesIO()
        workbook.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="formulas-bulk-import-template.xlsx"'
        return response
