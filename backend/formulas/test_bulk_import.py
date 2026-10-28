import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.authtoken.models import Token
from rest_framework.test import APIClient

from .models import TECH_DATA_AUTO_SECTION_LABEL, FormulaDefinition, Section, TechDataFieldDefinition

User = get_user_model()

URL = "/api/formulas/bulk-import/"

SECTIONS = [
    ["label", "sheet", "order"],
    ["9. Coupling Parts", "pricing_tool", 24],
    ["Extra Inputs", "tech_sheet", None],
]
FIELDS = [
    ["label", "key", "section", "type", "unit", "fixed_value", "auto", "expression"],
    ["Coupling Weight", "couplingWeight", "9. Coupling Parts", "number", "kg", "25", "No", None],
    ["Coupling Type", None, "9. Coupling Parts", "select", None, None, "No", None],
    ["Coupling Cost", None, "9. Coupling Parts", "number", "INR", None, "Yes", "couplingWeight * 120"],
]


def workbook_file(sections=None, fields=None, name="import.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in (("Sections", sections), ("Fields", fields)):
        if rows is None:
            continue
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class BulkImportTests(TestCase):
    def setUp(self):
        controlling = User.objects.create_user("controlling-bulk", password="controllingpass123", role="Controlling")
        sales = User.objects.create_user("sales-bulk", password="salespass123", role="Sales")
        self.client_ = APIClient()
        self.client_.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=controlling).key}")
        self.sales = APIClient()
        self.sales.credentials(HTTP_AUTHORIZATION=f"Token {Token.objects.create(user=sales).key}")

    def upload(self, sections=None, fields=None, dry_run=False, client=None, file=None):
        return (client or self.client_).post(
            URL,
            {"file": file or workbook_file(sections, fields), "dry_run": "true" if dry_run else "false"},
            format="multipart",
        )

    def test_template_has_both_sheets_with_the_agreed_columns(self):
        res = self.client_.get(URL + "template/")
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(io.BytesIO(res.content))
        self.assertEqual([c.value for c in wb["Sections"][1]], ["label", "sheet", "order"])
        self.assertEqual(
            [c.value for c in wb["Fields"][1]],
            ["label", "key", "section", "type", "unit", "fixed_value", "auto", "expression"],
        )
        # Example rows under the headings show how to fill in each sheet.
        self.assertEqual([c.value for c in wb["Sections"][2]], ["9. Coupling Parts", "pricing_tool", 24])
        self.assertEqual([wb["Fields"].cell(row=r, column=1).value for r in (2, 3, 4)], ["Coupling Weight", "Coupling Type", "Coupling Cost"])

    def test_template_examples_are_valid_rows(self):
        template = self.client_.get(URL + "template/").content
        res = self.upload(file=SimpleUploadedFile("template.xlsx", template), dry_run=True)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["errors"], [])
        self.assertEqual(res.data["summary"]["sections_created"], 2)
        self.assertEqual(res.data["summary"]["fields_created"], 3)

    def test_dry_run_reports_everything_but_saves_nothing(self):
        res = self.upload(SECTIONS, FIELDS, dry_run=True)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertFalse(res.data["saved"])
        self.assertEqual(res.data["summary"]["sections_created"], 2)
        self.assertEqual(res.data["summary"]["fields_created"], 3)
        self.assertEqual(res.data["errors"], [])
        self.assertFalse(Section.objects.filter(label="9. Coupling Parts").exists())
        self.assertFalse(TechDataFieldDefinition.objects.filter(key="couplingWeight").exists())

    def test_import_creates_sections_and_manual_select_and_auto_fields(self):
        res = self.upload(SECTIONS, FIELDS)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertTrue(res.data["saved"])

        coupling = Section.objects.get(label="9. Coupling Parts")
        self.assertEqual((coupling.tab, coupling.order), ("pricing_tool", 24))
        self.assertEqual(Section.objects.get(label="Extra Inputs").tab, "tech_sheet")

        weight = TechDataFieldDefinition.objects.get(key="couplingWeight")
        self.assertEqual((weight.field_type, weight.unit, weight.is_auto, weight.is_core), ("number", "kg", False, False))
        self.assertEqual(weight.fixed_value, "25")

        # A dropdown is imported without choices (no options column) — added later with Edit.
        coupling_type = TechDataFieldDefinition.objects.get(key="couplingType")  # key made from the label
        self.assertEqual((coupling_type.field_type, coupling_type.options), ("select", []))
        type_row = next(r for r in res.data["fields"] if r["key"] == "couplingType")
        self.assertIn("add them with Edit", type_row["message"])

        cost = TechDataFieldDefinition.objects.get(key="couplingCost")
        self.assertTrue(cost.is_auto)
        self.assertEqual(cost.formula.expression, "couplingWeight * 120")
        self.assertEqual(cost.formula.input_variables, ["couplingWeight"])
        self.assertEqual(cost.formula.section, TECH_DATA_AUTO_SECTION_LABEL)
        self.assertLess(weight.order, coupling_type.order)

    def test_auto_field_may_use_a_field_listed_after_it_in_the_file(self):
        fields = [FIELDS[0], FIELDS[3], FIELDS[1]]  # Coupling Cost before Coupling Weight
        res = self.upload(SECTIONS, fields)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertTrue(TechDataFieldDefinition.objects.get(key="couplingCost").is_auto)

    def test_existing_section_and_field_are_skipped_and_the_rest_imported(self):
        Section.objects.create(key="couplingParts", label="9. Coupling Parts", order=3, tab="tech_sheet")
        TechDataFieldDefinition.objects.create(key="couplingWeight", label="Old label", section="9. Coupling Parts", field_type="text")
        res = self.upload(SECTIONS, FIELDS)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["summary"]["sections_skipped"], 1)
        self.assertEqual(res.data["summary"]["fields_skipped"], 1)
        self.assertEqual(Section.objects.get(label="9. Coupling Parts").tab, "tech_sheet")  # unchanged
        self.assertEqual(TechDataFieldDefinition.objects.get(key="couplingWeight").label, "Old label")  # unchanged
        self.assertTrue(TechDataFieldDefinition.objects.filter(key="couplingCost").exists())

    def test_any_error_rolls_back_the_whole_file(self):
        fields = [row[:] for row in FIELDS]
        fields[1][5] = "heavy"  # fixed value of a Number field must be a number
        res = self.upload(SECTIONS, fields)
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data["errors"][0]["sheet"], "Fields")
        self.assertEqual(res.data["errors"][0]["row"], 2)
        self.assertIn("Must be a number", res.data["errors"][0]["message"])
        self.assertFalse(Section.objects.filter(label="9. Coupling Parts").exists())
        self.assertFalse(TechDataFieldDefinition.objects.filter(key="couplingWeight").exists())

    def test_row_errors_are_reported_with_their_row_number(self):
        fields = [
            FIELDS[0],
            ["Bad Formula", None, "9. Coupling Parts", "number", None, None, "Yes", "noSuchField * 2"],
            ["Uses Bad", None, "9. Coupling Parts", "number", None, None, "Yes", "badFormula + 1"],
            ["Fixed Auto", None, "9. Coupling Parts", "number", None, "5", "Yes", "1 + 1"],
            ["Wrong Section", None, "Not A Section", "text", None, None, "No", None],
        ]
        sections = [SECTIONS[0], SECTIONS[1], ["Bad Sheet", "somewhere", None]]
        res = self.upload(sections, fields, dry_run=True)
        self.assertEqual(res.status_code, 200)
        by_row = {(e["sheet"], e["row"]): e["message"] for e in res.data["errors"]}
        self.assertIn("pricing_tool", by_row[("Sections", 3)])
        self.assertIn("Unknown variable", by_row[("Fields", 2)])
        self.assertIn("row 2", by_row[("Fields", 3)])
        self.assertIn("Auto field can't have a fixed value", by_row[("Fields", 4)])
        self.assertIn("known section", by_row[("Fields", 5)])

    def test_select_fixed_value_becomes_its_choice_and_other_fields_can_copy_it(self):
        # The sheet from the user's test: a dropdown with a fixed value, an Auto field
        # copying it, and another Auto field copying that one.
        sections = [SECTIONS[0], ["Bulk Import Test1", "pricing_tool", None]]
        fields = [
            FIELDS[0],
            ["Bulk", "bulk", "Bulk Import Test1", "select", None, "SNL 3172", "No", None],
            ["Bulk2", "bulk2", "Bulk Import Test1", "number", "kg", None, "Yes", "bulk"],
            ["Bulk3", "bulk3", "Bulk Import Test1", "text", None, None, None, None],
            ["Bulk4", "bulk4", "Bulk Import Test1", "number", "kg", None, "yes", "bulk2"],
        ]
        res = self.upload(sections, fields)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["errors"], [])
        bulk = TechDataFieldDefinition.objects.get(key="bulk")
        self.assertEqual((bulk.field_type, bulk.options, bulk.fixed_value), ("select", ["SNL 3172"], "SNL 3172"))
        self.assertIn("one choice, 'SNL 3172'", next(r for r in res.data["fields"] if r["key"] == "bulk")["message"])
        self.assertEqual(TechDataFieldDefinition.objects.get(key="bulk2").formula.expression, "bulk")
        self.assertEqual(TechDataFieldDefinition.objects.get(key="bulk4").formula.expression, "bulk2")

    def test_text_auto_field_must_copy_a_single_field(self):
        fields = [FIELDS[0], ["Weight Text", None, "9. Coupling Parts", "text", None, None, "Yes", "couplingWeight * 2"]]
        res = self.upload(SECTIONS, fields, dry_run=True)
        self.assertIn("can only copy another field", res.data["errors"][0]["message"])

    def test_auto_key_clashing_with_an_existing_formula_is_an_error(self):
        FormulaDefinition.objects.create(key="couplingCost", label="x", section="S", expression="1")
        res = self.upload(SECTIONS, FIELDS, dry_run=True)
        self.assertIn("formula with the key 'couplingCost' already exists", res.data["errors"][0]["message"])

    def test_sales_cannot_import(self):
        self.assertEqual(self.upload(SECTIONS, FIELDS, client=self.sales).status_code, 403)

    def test_non_excel_file_is_rejected(self):
        res = self.upload(file=SimpleUploadedFile("fields.csv", b"label\nx", content_type="text/csv"))
        self.assertEqual(res.status_code, 400)
        self.assertIn(".xlsx", res.data["detail"])
